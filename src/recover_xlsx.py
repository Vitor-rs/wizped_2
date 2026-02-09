import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import os

# Paths
source_file = "docs/gerador_fichas_freq.xlsm"
target_file = "ficha_freq_gerador.xlsx"

# Sheet to Table Name mapping (derived from analysis)
tables_map = {
    "BD_Alunos": "Tbl_Alunos",
    "BD_Agenda": "Tbl_Agenda",
    "BD_Livros": "Tbl_Livros",
    "BD_Experiencia": "Tbl_Experiencia",
    "BD_Modalidades": "Tbl_Modalidades",
    "BD_Status": "Tbl_Status",
    "BD_Contrato": "Tbl_Contrato",
    "BD_Professores": "Tbl_Professores",
    "BD_Horarios": "Tbl_Horarios",
}


def recover_data():
    if not os.path.exists(source_file):
        print(f"Error: Source file '{source_file}' not found.")
        return

    print(f"Reading data from {source_file}...")

    # Needs openpyxl engine. Warning: might warn about extensions, but usually reads data.
    try:
        # Read all sheets
        xls = pd.read_excel(source_file, sheet_name=None, engine="openpyxl")
    except Exception as e:
        print(f"Pandas failed to read: {e}")
        # Fallback? attempt specific sheets?
        return

    print(f"Found {len(xls)} sheets: {list(xls.keys())}")

    # Create new writer
    with pd.ExcelWriter(target_file, engine="openpyxl") as writer:
        for sheet_name, df in xls.items():
            print(f"Processing sheet: {sheet_name} ({len(df)} rows)")

            # Write data to new sheet
            df.to_excel(writer, sheet_name=sheet_name, index=False)

            # Access the workbook to add tables
            worksheet = writer.sheets[sheet_name]

            # Formatting as Table (ListObject)
            if sheet_name in tables_map and not df.empty:
                table_name = tables_map[sheet_name]
                # Calculate range ref (e.g., A1:C10)
                # openpyxl uses 1-based indexing.
                # min_row=1 (header), max_row=len(df)+1
                # min_col=1, max_col=len(columns)
                from openpyxl.utils import get_column_letter

                max_col = len(df.columns)
                max_row = len(df) + 1
                ref = f"A1:{get_column_letter(max_col)}{max_row}"

                print(f"  Creating table '{table_name}' at {ref}")

                tab = Table(displayName=table_name, ref=ref)

                # Add a default style with striped rows
                style = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                tab.tableStyleInfo = style

                worksheet.add_table(tab)
            else:
                print(
                    f"  Skipping table creation for {sheet_name} (empty or not mapped)"
                )

    print(f"Successfully created '{target_file}'")


if __name__ == "__main__":
    recover_data()
